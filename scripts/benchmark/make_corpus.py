"""Builds the benchmark corpus from this repository's own text.

Real text files are copied. The Office and PDF files are generated from the repository's own prose
(clearly synthetic containers, real content), because there is no public set of real documents to ship.
Usage: python3 scripts/benchmark/make_corpus.py OUT_DIR
Needs: pip install python-docx openpyxl python-pptx reportlab
and Google Chrome for the browser-exported PDF (via playwright-core from node_modules).
"""
import csv, json, os, re, shutil, subprocess, sys
from collections import Counter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
out = sys.argv[1]
os.makedirs(out, exist_ok=True)

for src, dst in [("docs/AGENT_GUIDE.md", "agent-guide.md"), ("PRIVACY.md", "privacy.md"), ("ROADMAP.md", "roadmap.md"),
                 ("site/index.html", "site.html"), ("docs/examples/manifest.example.json", "manifest.json"),
                 ("package.json", "package.json"), ("scripts/deploy.mjs", "deploy.mjs"), ("src/app.ts", "app.ts")]:
    shutil.copy(os.path.join(ROOT, src), os.path.join(out, dst))

log = subprocess.check_output(["git", "log", "--format=%h|%an|%ad|%s", "--date=short"], cwd=ROOT, text=True).strip().splitlines()
commits = [line.split("|", 3) for line in log]
with open(os.path.join(out, "commits.csv"), "w", newline="") as f:
    writer = csv.writer(f); writer.writerow(["hash", "author", "date", "subject"]); writer.writerows(commits)

lock = json.load(open(os.path.join(ROOT, "package-lock.json")))
with open(os.path.join(out, "packages.csv"), "w", newline="") as f:
    writer = csv.writer(f); writer.writerow(["package", "version", "license", "resolved"])
    for key, value in lock["packages"].items():
        if key:
            writer.writerow([key.replace("node_modules/", ""), value.get("version", ""), value.get("license", ""), value.get("resolved", "")])

import docx
document = docx.Document()
for line in open(os.path.join(ROOT, "PRIVACY.md")).read().splitlines():
    if line.startswith("# "): document.add_heading(line[2:], 1)
    elif line.startswith("## "): document.add_heading(line[3:], 2)
    elif line.startswith("- "): document.add_paragraph(re.sub(r"[*`]", "", line[2:]), style="List Bullet")
    elif line.strip(): document.add_paragraph(re.sub(r"[*`]", "", line))
table = document.add_table(rows=1, cols=3)
for i, heading in enumerate(["Item", "Stored", "Deleted"]): table.rows[0].cells[i].text = heading
for row in [("Originals", "R2", "window end"), ("Projections", "R2", "window end"), ("File records", "D1", "window end")]:
    cells = table.add_row().cells
    for i, value in enumerate(row): cells[i].text = value
document.save(os.path.join(out, "privacy.docx"))

import openpyxl
book = openpyxl.Workbook(); sheet = book.active; sheet.title = "commits"; sheet.append(["hash", "author", "date", "subject"])
for row in commits: sheet.append(row)
by_author = book.create_sheet("by_author")
for author, count in Counter(row[1] for row in commits).items(): by_author.append([author, count])
book.save(os.path.join(out, "commits.xlsx"))

from pptx import Presentation
deck = Presentation()
for section in re.split(r"\n## ", open(os.path.join(ROOT, "README.md")).read())[1:8]:
    title, *body = section.split("\n", 1)
    slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = title.strip()
    lines = [re.sub(r"[*`|\[\]]", "", l).strip() for l in (body[0] if body else "").splitlines() if l.strip() and not l.startswith(("```", "|"))][:6]
    slide.placeholders[1].text = "\n".join(lines)
deck.save(os.path.join(out, "readme-deck.pptx"))

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
pdf = canvas.Canvas(os.path.join(out, "agent-guide.pdf"), pagesize=letter); y = 750
for line in open(os.path.join(ROOT, "docs/AGENT_GUIDE.md")).read().splitlines()[:140]:
    for chunk in [line[i:i + 95] for i in range(0, max(len(line), 1), 95)]:
        pdf.drawString(40, y, chunk); y -= 13
        if y < 50: pdf.showPage(); y = 750
pdf.save()

# A PDF exported from a browser, the way most people make one.
chrome = os.path.join(os.path.dirname(__file__), "chrome_pdf.mjs")
subprocess.check_call(["node", chrome, os.path.join(ROOT, "PRIVACY.md"), os.path.join(out, "privacy-chrome.pdf")])
print(sorted(os.listdir(out)))
