"""Compares extracted text with independent extractors, word by word.

Usage:
  EVAL_IN=DIR EVAL_OUT=OUT npx vitest run tests/extraction-dump.test.ts
  python3 scripts/benchmark/compare_extraction.py DIR OUT

For each file, recall is the share of the reference extractor's words that appear in ours, and
precision is the share of our words the reference also has. Words are lower-case letters and digits.
References: poppler's pdftotext for PDFs (best of its raw, default, and layout modes, since they
differ on word spacing); python-docx, python-pptx, and openpyxl for Office files; pandoc for EPUB and HTML. Precision is
lower for Office files by design, because we add headings, slide numbers, and metadata lines.
Needs: pdftotext and pandoc on PATH, and pip install python-docx python-pptx openpyxl
"""
import os, re, subprocess, sys
from collections import Counter

directory, out = sys.argv[1], sys.argv[2]
words = lambda text: Counter(re.findall(r"[a-z0-9]+", text.lower()))


def ref_docx(path):
    import docx
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    for section in document.sections:
        parts += [p.text for p in section.header.paragraphs] + [p.text for p in section.footer.paragraphs]
    return "\n".join(parts)


def ref_pptx(path):
    from pptx import Presentation
    parts = []
    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 6: walk(shape.shapes)
            if shape.has_text_frame: parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows: parts.extend(cell.text for cell in row.cells)
    for slide in Presentation(path).slides:
        walk(slide.shapes)
        if slide.has_notes_slide: parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def ref_xlsx(path):
    import datetime, openpyxl
    values, formulas = openpyxl.load_workbook(path, data_only=True), openpyxl.load_workbook(path)
    def show(value):
        if isinstance(value, datetime.datetime):
            return value.strftime("%Y-%m-%d") if value.time() == datetime.time() else value.strftime("%Y-%m-%d %H:%M")
        if isinstance(value, bool): return "TRUE" if value else "FALSE"
        if isinstance(value, float): return repr(float(f"{value:.15g}"))
        return str(value)
    parts = []
    for sheet in formulas.worksheets:
        cached = values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                value = cached[cell.coordinate].value if cached[cell.coordinate].value is not None else cell.value
                if value is not None: parts.append(show(value))
    return "\n".join(parts)


def best_pdf_reference(path, ours):
    best = None
    for mode in (["-raw"], [], ["-layout"]):
        text = subprocess.run(["pdftotext", *mode, path, "-"], capture_output=True, text=True).stdout
        ref = words(text)
        score = sum((ref & ours).values()) / max(sum(ref.values()), 1)
        if best is None or score > best[0]: best = (score, ref)
    return best[1]


def ref_pandoc(path, source):
    return subprocess.run(["pandoc", "-f", source, "-t", "plain", "--wrap=none", path], capture_output=True, text=True).stdout


rows = []
for name in sorted(os.listdir(directory)):
    ext = name.rsplit(".", 1)[-1].lower()
    if ext not in ("pdf", "docx", "pptx", "xlsx", "epub", "html", "htm"): continue
    ours = words(open(os.path.join(out, name + ".txt")).read())
    path = os.path.join(directory, name)
    if ext == "pdf": ref = best_pdf_reference(path, ours)
    elif ext == "epub": ref = words(ref_pandoc(path, "epub"))
    elif ext in ("html", "htm"): ref = words(ref_pandoc(path, "html"))
    else: ref = words({"docx": ref_docx, "pptx": ref_pptx, "xlsx": ref_xlsx}[ext](path))
    total = sum(ref.values())
    shared = sum((ref & ours).values())
    rows.append((name, total, shared / max(total, 1), shared / max(sum(ours.values()), 1)))

print(f'{"file":46}{"ref words":>10}{"recall":>8}{"precision":>10}')
for row in sorted(rows, key=lambda r: r[2]): print(f"{row[0][:45]:46}{row[1]:>10}{row[2]:>8.3f}{row[3]:>10.3f}")
weight = sum(r[1] for r in rows)
print(f"\nfiles: {len(rows)}  median recall {sorted(r[2] for r in rows)[len(rows) // 2]:.3f}  word-weighted recall {sum(r[1] * r[2] for r in rows) / max(weight, 1):.3f}  below 0.90: {sum(r[2] < 0.9 for r in rows)}")
